"""보고서 HTML의 내용을 편집 가능한 PowerPoint 요소로 변환."""
from pathlib import Path
from bs4 import BeautifulSoup
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

HERE = Path(__file__).resolve().parent
BLUE = '1C50B3'
MAGENTA = 'BE1551'
FONT = '맑은 고딕'
prs = Presentation()
prs.slide_width = Inches(13.333333)
prs.slide_height = Inches(7.5)


def pos(px):
    return Inches(px / 96)


def rect(slide, x, y, w, h, color, rounded=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        pos(x), pos(y), pos(w), pos(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()
    if rounded:
        shape.adjustments[0] = .05
    return shape


def text(slide, x, y, w, h, paragraphs, size=22, bold=False, color='080808',
         align=PP_ALIGN.LEFT, font=FONT, spacing=12):
    shape = slide.shapes.add_textbox(pos(x), pos(y), pos(w), pos(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = pos(2)
    tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.alignment = align
        p.font.name = font
        p.font.size = Pt(size * .75)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(spacing * .75)
        p.line_spacing = 1.2
    return shape


def picture(slide, src, x, y, w, h):
    path = (HERE / src).resolve()
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    slide.shapes.add_picture(str(path), pos(x + (w-pw)/2), pos(y+(h-ph)/2),
                             width=pos(pw), height=pos(ph))


def panel(slide, aside, x, y, w, h, size=22):
    rect(slide, x, y, w, h, BLUE, True)
    rect(slide, x+9, y+54, w-18, h-64, 'FFFFFF')
    text(slide, x+22, y+12, w-44, 38, [aside.h3.get_text(' ', strip=True)],
         size=23, bold=True, color='FFFFFF')
    paragraphs = [p.get_text(' ', strip=True) for p in aside.select('p')]
    text(slide, x+25, y+68, w-50, h-83, paragraphs, size=size, bold=True)


def native_table(slide, node, x, y, w, h):
    rows = node.find_all('tr')
    cols = len(rows[0].find_all(['th', 'td']))
    shape = slide.shapes.add_table(len(rows), cols, pos(x), pos(y), pos(w), pos(h))
    table = shape.table
    for r, row in enumerate(rows):
        for c, cell_node in enumerate(row.find_all(['th', 'td'])):
            cell = table.cell(r, c)
            cell.text = cell_node.get_text(strip=True)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(
                BLUE if r == 0 else ('EEF3FB' if r % 2 == 0 else 'FFFFFF'))
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER
                p.font.name = FONT
                p.font.size = Pt(16.5)
                p.font.bold = r == 0
                p.font.color.rgb = RGBColor.from_string('FFFFFF' if r == 0 else '080808')


soup = BeautifulSoup((HERE / '강화학습_week4_2021042040.html').read_text(), 'html.parser')
for index, node in enumerate(soup.select('section.slide'), 1):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if 'cover' in node.get('class', []):
        rect(slide, 0, 0, 1280, 596, BLUE)
        rect(slide, 0, 81, 1280, 2, 'FFFFFF')
        text(slide, 18, 12, 300, 65, ['CBNU'], size=54, bold=True, color=MAGENTA, font='Arial')
        text(slide, 48, 150, 1100, 220, ['강화학습', 'Week4 HW'], size=72,
             bold=True, color='FFFFFF', spacing=10)
        text(slide, 50, 390, 1000, 60, ['Quiz 1 · Quiz 2'], size=29, color='FFFFFF')
        text(slide, 28, 642, 1150, 50, [node.select_one('.identity').get_text()], size=25, bold=True)
        continue
    rect(slide, 0, 0, 1280, 82, BLUE)
    text(slide, 27, 17, 1220, 55, [node.header.get_text()], size=40, bold=True, color='FFFFFF')
    text(slide, 24, 696, 100, 24, [f'{index:02d}'], size=15, color='777777')
    text(slide, 1150, 685, 128, 35, ['CBNU'], size=30, bold=True, color=MAGENTA, font='Arial',
         align=PP_ALIGN.RIGHT)
    if node.select_one('.contents'):
        y = 140
        for item in node.select_one('.contents').find_all(['h2', 'p'], recursive=False):
            is_heading = item.name == 'h2'
            paragraphs = item.get_text('\n', strip=True).split('\n')
            height = 48 if is_heading else 110
            text(slide, 475 if is_heading else 500, y, 740, height, paragraphs,
                 size=30 if is_heading else 24, bold=is_heading, spacing=2)
            y += height + 12
    elif node.select_one('.divider'):
        parts = node.select_one('.divider').get_text('\n', strip=True).split('\n')
        text(slide, 30, 290, 1220, 105, [parts[0]], size=65, bold=True, align=PP_ALIGN.CENTER)
        text(slide, 30, 410, 1220, 80, [parts[1]], size=37, align=PP_ALIGN.CENTER)
    elif node.select_one('.split'):
        visual = node.select_one('.visual')
        if visual.img:
            picture(slide, visual.img['src'], 36, 130, 560, 510)
        else:
            source = visual.get_text().strip()
            rect(slide, 36, 135, 560, 510, '272822')
            text(slide, 50, 150, 533, 483, source.splitlines(), size=15,
                 color='F8F8F2', font='Consolas', spacing=0)
        panel(slide, node.aside, 642, 133, 600, 510, size=21)
    elif node.select_one('.full'):
        full = node.select_one('.full')
        native_table(slide, full.table, 50, 122, 1180, 318)
        panel(slide, full.aside, 50, 467, 1180, 210, size=22)
    elif node.select_one('.pair'):
        for col, item in enumerate(node.select_one('.pair').find_all('div', recursive=False)):
            x = 60 + col*610
            text(slide, x, 108, 550, 45, [item.h3.get_text()], size=25, bold=True, align=PP_ALIGN.CENTER)
            picture(slide, item.img['src'], x, 150, 550, 460)
        text(slide, 50, 642, 1180, 45, [node.select_one('.caption').get_text()],
             size=24, align=PP_ALIGN.CENTER)

target = HERE / '강화학습_week4_2021042040.pptx'
prs.save(target)
print(f'{len(prs.slides)} editable slides: {target}')
